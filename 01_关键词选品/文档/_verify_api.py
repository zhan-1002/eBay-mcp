# -*- coding: utf-8 -*-
"""核对 API 模式产物：条数 / item specifics / 类目 / 图片 / 描述 / 广告位 / 失败数。"""
import glob
import json
import os

import openpyxl

OUT = r"C:\Users\admin\Desktop\新建文件夹 (2)\紫鸟ebay\01_关键词选品\输出"
J = sorted(glob.glob(os.path.join(OUT, "api_uk_wireless_earbuds_*.json")))[-1]
print("核对文件: %s" % os.path.basename(J))
d = json.load(open(J, encoding="utf-8"))
its = d["items"]
n = len(its)


def c(pred):
    return sum(1 for it in its if pred(it))


print("\n=== 采集完整性 ===")
print("  条数                      : %d" % n)
print("  唯一 item_id              : %d" % len({it.get("item_id") for it in its}))
print("  legacy_item_id 非空       : %d" % c(lambda x: x.get("legacy_item_id")))
print("  item_url 非空             : %d" % c(lambda x: x.get("item_url")))
print("  详情成功(detail_ok)       : %d" % c(lambda x: x.get("detail_ok")))
print("  详情失败                  : %d %s"
      % (c(lambda x: x.get("detail_error")),
         [x["item_id"][:22] for x in its if x.get("detail_error")][:3]))

print("\n=== item specifics（本次核心）===")
print("  有 item_specifics         : %d / %d = %.1f%%"
      % (c(lambda x: x.get("item_specifics")), n, 100.0 * c(lambda x: x.get("item_specifics")) / n))
cnts = [len(x.get("item_specifics") or []) for x in its]
print("  每条属性数 min/中位/max    : %d / %d / %d"
      % (min(cnts), sorted(cnts)[len(cnts) // 2], max(cnts)))
from collections import Counter
namefreq = Counter()
for x in its:
    for s in x.get("item_specifics") or []:
        namefreq[s["name"]] += 1
print("  属性名种类                : %d" % len(namefreq))
print("  出现最多 8 个             : %s" % namefreq.most_common(8))
NEED = ["Brand", "Connectivity", "Model", "Colour", "Type"]
print("  刊登必填项覆盖:")
for k in NEED:
    print("     %-14s %d/%d = %.1f%%"
          % (k, sum(1 for x in its if any(s["name"].lower() == k.lower()
                                          for s in x.get("item_specifics") or [])),
             n, 100.0 * sum(1 for x in its if any(s["name"].lower() == k.lower()
                                                  for s in x.get("item_specifics") or [])) / n))

print("\n=== 类目 ===")
print("  leaf_category_ids 非空    : %d" % c(lambda x: x.get("leaf_category_ids")))
print("  category_path 非空        : %d" % c(lambda x: x.get("category_path")))
print("  有第二类目               : %d %s"
      % (c(lambda x: x.get("has_secondary_category")),
         [(x["position"], x["leaf_category_ids"]) for x in its if x.get("has_secondary_category")][:4]))
print("  末级类目分布 top5         : %s" % Counter(
    (x.get("leaf_category_ids") or ["-"])[0] for x in its).most_common(5))

print("\n=== 其它字段 ===")
print("  价格非空                  : %d" % c(lambda x: x.get("price") is not None))
print("  广告位(priorityListing)   : %d" % c(lambda x: x.get("is_sponsored")))
print("  图片 URL 非空             : %d" % c(lambda x: x.get("image_url")))
print("  描述非空                  : %d" % c(lambda x: x.get("description_len")))
print("  卖家非空                  : %d" % c(lambda x: x.get("seller")))
print("  商品所在地非空            : %d" % c(lambda x: x.get("item_location")))

print("\n=== Excel 子表 ===")
X = J[:-5] + ".xlsx"
wb = openpyxl.load_workbook(X)
print("  共 %d 个子表: %s" % (len(wb.sheetnames), wb.sheetnames))
hdr = [cc.value for cc in wb["采集明细"][1]]
print("\n  采集明细列（%d 列）:" % len(hdr))
print("  %s" % hdr)
print("\n  行数: %d" % wb["采集明细"].max_row)

print("\n=== 前 3 条样例 ===")
for it in its[:3]:
    print("\n  #%d %s" % (it["position"], (it["title"] or "")[:64]))
    print("     id=%s | £%s | %s | 广告=%s | 图片 %d 张"
          % (it["legacy_item_id"], it["price"], it["item_location"], it["is_sponsored"],
             it.get("image_count") or 0))
    print("     类目: %s | %s" % (it["leaf_category_ids"], it.get("category_path")))
    print("     属性 %d 项: %s" % (len(it["item_specifics"]),
                                ", ".join("%s=%s" % (s["name"], str(s["value"])[:16])
                                          for s in (it["item_specifics"] or [])[:6])))
    print("     描述 %d 字符" % (it.get("description_len") or 0))
