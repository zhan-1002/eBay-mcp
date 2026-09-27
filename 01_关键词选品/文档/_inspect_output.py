# -*- coding: utf-8 -*-
"""完整体检 API 模式产物：JSON 字段 + Excel 各子表内容。"""
import glob
import json
import os

import openpyxl

OUT = r"C:\Users\admin\Desktop\新建文件夹 (2)\紫鸟ebay\01_关键词选品\输出"
TARGET = sorted(glob.glob(os.path.join(OUT, "api_uk_wireless_earbuds_*.json")))[-2]  # 倒数第二个 = 完整 120 条
J = TARGET
X = J[:-5] + ".xlsx"
print("=" * 88)
print("体检对象")
print("=" * 88)
print("  JSON: %s" % os.path.basename(J))
print("  XLSX: %s" % os.path.basename(X))
print("  大小: %.0f KB / %.0f KB" % (os.path.getsize(J) / 1024, os.path.getsize(X) / 1024))

d = json.load(open(J, encoding="utf-8"))
its = d["items"]
n = len(its)
print("  条数: %d | 关键词: %s | 站点: %s" % (n, d.get("keyword"), d.get("site")))

print()
print("=" * 88)
print("① Excel 子表清单与规模")
print("=" * 88)
wb = openpyxl.load_workbook(X)
for ws in wb:
    print("  %-18s %4d 行 x %2d 列" % (ws.title, ws.max_row, ws.max_column))

print()
print("=" * 88)
print("② 市场概况（AI 直接读这张）")
print("=" * 88)
for r in wb["市场概况"].iter_rows(values_only=True):
    if any(x is not None for x in r):
        print("  %-16s | %s" % (r[0], r[1]))

print()
print("=" * 88)
print("③ 价格分析（含占比）")
print("=" * 88)
for r in wb["价格分析"].iter_rows(values_only=True):
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c) for c in r[:7]))

print()
print("=" * 88)
print("④ 品牌壁垒（前 12 行）")
print("=" * 88)
for i, r in enumerate(wb["品牌壁垒"].iter_rows(values_only=True), 1):
    if i > 12:
        break
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c)[:28] for c in r[:11]))

print()
print("=" * 88)
print("⑤ 关键词与标题（前 14 行 + 末尾推荐标题）")
print("=" * 88)
ws = wb["关键词与标题"]
rows = list(ws.iter_rows(values_only=True))
for r in rows[:14]:
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c)[:22] for c in r[:10]))
print("  ...")
for r in rows[-6:]:
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c)[:70] for c in r[:3]))

print()
print("=" * 88)
print("⑥ 采集明细列（%d 列）与内容完整度" % wb["采集明细"].max_column)
print("=" * 88)
hdr = [c.value for c in wb["采集明细"][1]]
print("  列: %s" % hdr)


def c(pred):
    return sum(1 for it in its if pred(it))


print()
print("  item_specifics 非空   : %d/%d" % (c(lambda x: x.get("item_specifics")), n))
print("  description 非空      : %d/%d" % (c(lambda x: x.get("description")), n))
print("  images 数组非空       : %d/%d" % (c(lambda x: x.get("images")), n))
print("  category_path 非空    : %d/%d" % (c(lambda x: x.get("category_path")), n))
print("  第二类目              : %d 条" % c(lambda x: x.get("has_secondary_category")))
