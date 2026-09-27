# -*- coding: utf-8 -*-
"""校验 Excel 里那 30 条推荐标题是否真的满足硬约束（临时诊断脚本）。

同时回答两个问题：
  1. 是不是 LLM 版（而不是规则版兜底）—— 规则版长度只有 50 上下、且是"关键词堆叠"形态
  2. 逐条跑 llm_titles.validate_titles（与流水线同一份校验代码）+ 品牌词残留检查
"""
import io
import json
import os
import sys
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.normpath(os.path.join(HERE, "..", "脚本"))
if SCRIPTS not in sys.path:
    sys.path.insert(0, SCRIPTS)

from openpyxl import load_workbook  # noqa: E402
import llm_titles  # noqa: E402

xlsx = sys.argv[1]
json_path = os.path.splitext(xlsx)[0] + ".json"

wb = load_workbook(xlsx, data_only=True)
ws = wb["关键词与标题"]

titles = []
in_block = False
for row in ws.iter_rows(values_only=True):
    if row and isinstance(row[0], str) and row[0].startswith("五、推荐商品标题"):
        in_block = True
        continue
    if in_block:
        if row and isinstance(row[0], str) and row[0].startswith("六、"):
            break
        # 表头行是 ("序号","标题","字符数")，要跳过，别当成一条标题
        if row and isinstance(row[0], int) and row[1]:
            titles.append(str(row[1]).strip())

print("=" * 74)
print("标题校验：%s" % os.path.basename(xlsx))
print("=" * 74)
print("标题条数: %d" % len(titles))

lens = sorted(len(t) for t in titles)
print("字符数: 最小 %d / 中位 %d / 最大 %d ｜ ≥70 的 %d 条 ｜ =80 的 %d 条"
      % (lens[0], lens[len(lens) // 2], lens[-1],
         sum(1 for x in lens if x >= 70), sum(1 for x in lens if x == 80)))

# 规则版兜底特征：长度普遍 <60 且以关键词开头
rule_like = sum(1 for t in titles if len(t) < 60)
print("长度 <60 的条数: %d %s" % (rule_like, "→ ⚠️ 疑似规则版兜底" if rule_like > 5 else "→ 是 LLM 版"))

# 与流水线同一份机械校验
with io.open(json_path, encoding="utf-8") as f:
    report = json.load(f)
market = report.get("market") or {}
meta = report.get("llm_titles") or {}
print("LLM 元信息: used=%s %s" % (meta.get("used"), meta.get("meta") or meta.get("reason")))

ok, problems = llm_titles.validate_titles(titles, [], [""], report.get("keyword") or "")
print("机械校验通过: %d/%d ｜ 问题 %d 条" % (len(ok), len(titles), len(problems)))
for p in problems[:10]:
    print("   - %s" % p)

# 逐条查每个约束（自己再算一遍，不依赖上面的函数）
bad_len = [t for t in titles if len(t) > 80]
dup_word = []
for t in titles:
    toks = __import__("re").findall(r"[a-z0-9&\-\.]+", t.lower())
    if len(toks) != len(set(toks)):
        dup_word.append(t)
punct = [t for t in titles if __import__("re").search(r"[^\w\s&\-\.]", t)]
dup_title = [t for t, c in Counter(t.lower() for t in titles).items() if c > 1]
print("超 80 字符: %d ｜ 内部重复词: %d ｜ 含标点: %d ｜ 互为重复: %d"
      % (len(bad_len), len(dup_word), len(punct), len(dup_title)))

# 品牌词残留：用真实 brand 字段里的品牌 + 标题品牌词表两边查
st = ((market.get("brand") or {}).get("structured") or {})
real_brands = [r.get("品牌") for r in (st.get("brand_rows") or [])]
tm = [r.get("标题品牌词") for r in ((market.get("brand") or {}).get("title_mentions") or [])]
print("-" * 74)
print("数据里出现的真品牌（brand 字段，前 20）: %s" % "、".join(str(b) for b in real_brands[:20]))
print("标题品牌词表（前 20）: %s" % "、".join(str(b) for b in tm[:20]))
hits = {}
for b in {str(x).lower() for x in real_brands + tm if x}:
    if len(b) < 3:
        continue
    for t in titles:
        if b in t.lower():
            hits.setdefault(b, []).append(t)
if hits:
    print("⚠️ 标题里命中品牌词表: %s" % "、".join("%s(%d 条)" % (k, len(v)) for k, v in hits.items()))
    for k, v in list(hits.items())[:3]:
        print("   %s → %s" % (k, v[0]))
else:
    print("标题里未命中任何品牌词表条目")

print("-" * 74)
print("30 条标题原文：")
for i, t in enumerate(titles, 1):
    print("  %2d. [%2d] %s" % (i, len(t), t))
