# -*- coding: utf-8 -*-
"""核对产物：Excel 各 sheet + JSON 字段完整性（取输出目录里最新一次运行）。"""
import glob
import json
import os

import openpyxl

OUT = r"C:\Users\admin\Desktop\新建文件夹 (2)\紫鸟ebay\01_关键词选品\输出"
XLSX = sorted(glob.glob(os.path.join(OUT, "haihu_8075_uk_wireless_earbuds_*.xlsx")))[-1]
JSON = XLSX[:-5] + ".json"
print("核对文件: %s" % os.path.basename(XLSX))

wb = openpyxl.load_workbook(XLSX)
print("\nsheets (%d): %s" % (len(wb.sheetnames), wb.sheetnames))
for ws in wb:
    print("  %-16s %d 行 x %d 列" % (ws.title, ws.max_row, ws.max_column))

d = json.load(open(JSON, encoding="utf-8"))
its = d["items"]
top10 = its[:10]

print("\n=== 核对项 ===")
print("1) 条数                 : %d" % len(its))
print("2) 广告位               : %d 条 (%.0f%%)" % (d["ad_count"], 100.0 * d["ad_count"] / len(its)))
print("3) 末级类目ID 覆盖       : %d / %d" % (sum(1 for it in its if it.get("leaf_category_ids")), len(its)))
print("4) 价格字段             : price/price_max/price_is_range 全在 = %s"
      % all(k in its[0] for k in ("price", "price_max", "price_is_range")))
print("   价格 min/p50/max     : %s / %s / %s %s"
      % (d["price"]["min"], d["price"]["p50"], d["price"]["max"], d["price"]["currency"]))
print("5) 前 10 条 specifics    : %s" % [len(it.get("item_specifics") or []) for it in top10])
print("   前 10 条里有 specifics : %d / 10" % sum(1 for it in top10 if it.get("item_specifics")))
print("   全文有 specifics 的    : %d / %d" % (sum(1 for it in its if it.get("item_specifics")), len(its)))
print("   属性种类              : %d 种" % len(d["specifics"]))
print("   属性名（前 12）        : %s" % [s["name"] for s in d["specifics"][:12]])
print("6) title 带 a11y 噪音    : %d 条"
      % sum(1 for it in its if "opens in a new window" in (it["title"] or "").lower()))
print("7) 推荐 title            : %d 条 ; 长度 %d~%d ; 含逗号 %d 条"
      % (len(d["recommended_titles"]),
         min(len(t) for t in d["recommended_titles"]),
         max(len(t) for t in d["recommended_titles"]),
         sum(1 for t in d["recommended_titles"] if "," in t)))
print("8) 第二类目             : %s" % [it["position"] for it in its if it.get("has_secondary_category")])
print("9) categories 非空       : %d 条" % sum(1 for it in its if it.get("categories")))
print("10) 类目名非空           : %d 条" % sum(1 for it in its if it.get("leaf_category_name")))

print("\n=== 前 10 条属性对照（看有没有 Brand/Type/Model/Colour） ===")
KEY = ["brand", "type", "model", "colour", "condition"]
for it in top10:
    low = [s["name"].lower() for s in it["item_specifics"]]
    have = [k for k in KEY if k in low]
    miss = [k for k in KEY if k not in low]
    print("  #%-3d %-44s 有=%-24s 缺=%s"
          % (it["position"], (it["title"] or "")[:44], ",".join(have), ",".join(miss)))

print("\n=== 采集明细表头 ===")
print([c.value for c in wb["采集明细"][1]])
