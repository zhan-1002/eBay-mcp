# -*- coding: utf-8 -*-
"""检查 5 表结构产物的内容质量（含新的品牌壁垒口径）。"""
import glob
import json
import os

import openpyxl

OUT = r"C:\Users\admin\Desktop\新建文件夹 (2)\紫鸟ebay\01_关键词选品\输出"
J = sorted(glob.glob(os.path.join(OUT, "api_uk_wireless_earbuds_*.json")))[-1]
X = J[:-5] + ".xlsx"
d = json.load(open(J, encoding="utf-8"))
its = d["items"]
print("检查: %s（%d 条）" % (os.path.basename(X), len(its)))

wb = openpyxl.load_workbook(X)
print("\n=== 子表（应为 5 张）===")
for ws in wb:
    print("  %-16s %4d 行 x %2d 列" % (ws.title, ws.max_row, ws.max_column))

print("\n" + "=" * 84)
print("① 市场概况与价格")
print("=" * 84)
for r in wb["市场概况与价格"].iter_rows(values_only=True):
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c)[:60] for c in r[:4]))

print("\n" + "=" * 84)
print("② 品牌壁垒 —— 一、结论（主口径）")
print("=" * 84)
ws = wb["品牌壁垒"]
rows = list(ws.iter_rows(values_only=True))
for r in rows[:12]:
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c)[:62] for c in r[:2]))

print("\n=== 二、各价格带真品牌占位率 ===")
for i, r in enumerate(rows):
    if r[0] and "二、" in str(r[0]):
        for r2 in rows[i:i + 10]:
            if any(x is not None for x in r2):
                print("  " + " | ".join("" if c is None else str(c)[:22] for c in r2[:5]))
        break

print("\n=== 三、各品牌真实份额（官方 brand 字段）===")
for i, r in enumerate(rows):
    if r[0] and "三、" in str(r[0]):
        for r2 in rows[i:i + 14]:
            if any(x is not None for x in r2):
                print("  " + " | ".join("" if c is None else str(c)[:24] for c in r2[:9]))
        break

print("\n=== 五、辅助：标题品牌词提及率（应带警示语）===")
for i, r in enumerate(rows):
    if r[0] and "五、" in str(r[0]):
        for r2 in rows[i:i + 8]:
            if any(x is not None for x in r2):
                print("  " + " | ".join("" if c is None else str(c)[:26] for c in r2[:6]))
        break

print("\n" + "=" * 84)
print("③ 关键词与标题（含模板/推荐标题/策略建议）")
print("=" * 84)
ws = wb["关键词与标题"]
for r in list(ws.iter_rows(values_only=True))[:6]:
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c)[:20] for c in r[:10]))
print("  ...")
seg = [r for r in ws.iter_rows(values_only=True) if r[0] and "五、" in str(r[0])]
print("  （末尾）%s" % (" | ".join(str(c)[:60] for c in seg[0] if c is not None) if seg else "无"))

print("\n" + "=" * 84)
print("④ 基础统计")
print("=" * 84)
for r in list(wb["基础统计"].iter_rows(values_only=True))[:6]:
    if any(x is not None for x in r):
        print("  " + " | ".join("" if c is None else str(c)[:30] for c in r[:4]))

print("\n" + "=" * 84)
print("⑤ 采集明细")
print("=" * 84)
print("  %d 行 x %d 列" % (wb["采集明细"].max_row, wb["采集明细"].max_column))

print("\n" + "=" * 84)
print("⑥ 数据完整度复核")
print("=" * 84)
n = len(its)


def c(p):
    return sum(1 for x in its if p(x))


print("  item specifics 非空 : %d/%d" % (c(lambda x: x.get("item_specifics")), n))
print("  brand_field 非空    : %d/%d" % (c(lambda x: x.get("brand_field")), n))
print("  描述非空            : %d/%d" % (c(lambda x: x.get("description")), n))
print("  category_path 非空  : %d/%d" % (c(lambda x: x.get("category_path")), n))
print("  第二类目            : %d 条" % c(lambda x: x.get("has_secondary_category")))
