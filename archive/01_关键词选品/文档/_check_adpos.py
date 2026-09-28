# -*- coding: utf-8 -*-
"""检查广告位位置分布两段是否写进 Excel（临时诊断脚本）。"""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=True)
print("子表:", wb.sheetnames)
ws = wb["市场概况与价格"]
for row in ws.iter_rows(values_only=True):
    vals = ["" if v is None else str(v) for v in row]
    line = " | ".join(vals).rstrip(" |")
    if line.strip():
        print(line)
