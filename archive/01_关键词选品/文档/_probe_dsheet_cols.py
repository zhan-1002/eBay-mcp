# -*- coding: utf-8 -*-
"""看 Soldeazy dsheet 导出表的真实列结构（含单品详情需要哪些列）。"""
import os

import openpyxl

for path in (r"D:\PythonProject\ebay\soldeazy\download\export_dsheet_UK.xlsx",
             r"D:\PythonProject\ebay\soldeazy\download\export_dsheet_DE.xlsx"):
    if not os.path.isfile(path):
        print("不存在: %s" % path)
        continue
    print("=" * 78)
    print("文件: %s  (%.0f KB)" % (path, os.path.getsize(path) / 1024))
    print("=" * 78)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb:
        print("--- sheet: %s  (%d 行 x %d 列) ---" % (ws.title, ws.max_row, ws.max_column))
        rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
        if not rows:
            continue
        hdr = rows[0]
        print("表头（列号: 名称）:")
        for i, h in enumerate(hdr, 1):
            if h is not None and str(h).strip():
                print("   %2d. %s" % (i, str(h)[:40]))
        print("\n前 3 行数据样例:")
        for r in rows[1:4]:
            cells = ["%s=%s" % (str(hdr[i])[:14], str(v)[:22])
                     for i, v in enumerate(r) if v is not None and i < len(hdr)]
            print("   %s" % " | ".join(cells[:12]))
    print()
