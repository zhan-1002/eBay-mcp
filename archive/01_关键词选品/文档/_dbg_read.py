# -*- coding: utf-8 -*-
"""直接读产物确认各段数据行是否存在（不用 iter_rows 的猜测）。"""
import glob
import os

import openpyxl

OUT = r"C:\Users\admin\Desktop\新建文件夹 (2)\紫鸟ebay\01_关键词选品\输出"
p = sorted(glob.glob(os.path.join(OUT, "json_uk_wireless_earbuds_*.xlsx")))[-1]
print("文件: %s" % os.path.basename(p))
wb = openpyxl.load_workbook(p)

for name in ("价格分析", "品牌壁垒", "关键词与标题"):
    ws = wb[name]
    print()
    print("=== %s: dims=%s max_row=%d ===" % (name, ws.dimensions, ws.max_row))
    for r in range(1, min(ws.max_row, 15) + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        d = ws.cell(row=r, column=4).value
        print("   r%-3d A=%-24r B=%-14r D=%r" % (r, a, b, d))
