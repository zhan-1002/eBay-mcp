# -*- coding: utf-8 -*-
"""检查「关键词与标题」子表里的 eBay 平台建议是否带上广告位位置分布（临时诊断脚本）。"""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=True)
ws = wb["关键词与标题"]
hit = False
for row in ws.iter_rows(values_only=True):
    vals = ["" if v is None else str(v) for v in row]
    line = " | ".join(vals).rstrip(" |")
    if "广告位" in line or "平台建议" in line:
        print(line)
        hit = True
if not hit:
    print("未找到广告位相关内容")
